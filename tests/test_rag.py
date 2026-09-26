import asyncio
import tempfile
import threading
import unittest
from pathlib import Path

import numpy as np
from types import SimpleNamespace

from src.core.andrew import AndrewCore
from src.core.llm import LLM
from src.core.memory import Memory
from src.core.rag import (
    KnowledgeBase,
    _Chunk,
    _Document,
    _query_texts,
    _table_heading,
    _table_rows,
)


class KnowledgeBaseTests(unittest.IsolatedAsyncioTestCase):
    async def test_retrieves_relevant_chunks_with_citations(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "rockets.txt").write_text(
                "The orbital launch sequence starts with propellant loading."
            )
            (root / "gardening.txt").write_text(
                "Tomatoes need regular watering and direct sunlight."
            )
            knowledge = KnowledgeBase(root, watch=False)
            knowledge.start()
            await knowledge._refresh_task

            context = await knowledge.retrieve("How does the orbital launch sequence start?")
            knowledge.stop()

        self.assertIn("[source: rockets.txt — rockets]", context)
        self.assertIn("propellant loading", context)
        self.assertNotIn("Tomatoes", context)

    async def test_page_window_links_an_entity_to_a_later_fact(self):
        document = _Document(
            "manual.pdf",
            (
                _Chunk("manual.pdf", "Overview", 1, "Orion systems division"),
                _Chunk("manual.pdf", "Overview", 2, "Application requirements"),
                _Chunk("manual.pdf", "Overview", 3, "Evaluation criteria"),
                _Chunk("manual.pdf", "Overview", 4, "Results are published at https://example.test/orion"),
                _Chunk("manual.pdf", "Overview", 10, "Other results are published elsewhere"),
            ),
        )
        class Reranker:
            def __init__(self):
                self.called = False

            def predict(self, pairs, **kwargs):
                self.called = True
                return [10.0 if "example.test/orion" in passage else 0.0 for _, passage in pairs]

        with tempfile.TemporaryDirectory() as directory:
            reranker = Reranker()
            knowledge = KnowledgeBase(Path(directory), watch=False, reranker_model="test")
            knowledge._reranker = reranker
            context = knowledge._format_context(
                "Where are the Orion systems results published? Do not use external search.",
                knowledge._build_snapshot([document]),
            )

        self.assertTrue(reranker.called)
        self.assertIn("https://example.test/orion", context)
        self.assertLessEqual(context.count("[source:"), 4)

    def test_dense_retrieval_recovers_lexically_unmatched_chunk(self):
        documents = [
            _Document("target.txt", (
                _Chunk("target.txt", "Guide", None, "vehicle service handbook"),
            )),
            _Document("other.txt", (
                _Chunk("other.txt", "Guide", None, "kitchen recipe collection"),
            )),
        ]
        vectors = {
            "target.txt": np.array([[1.0, 0.0]], dtype=np.float32),
            "other.txt": np.array([[0.0, 1.0]], dtype=np.float32),
        }

        class Embedder:
            def encode(self, texts, **kwargs):
                return np.array([[1.0, 0.0] for _ in texts], dtype=np.float32)

        with tempfile.TemporaryDirectory() as directory:
            knowledge = KnowledgeBase(
                Path(directory),
                watch=False,
                top_chunks=1,
                embedding_model="test",
            )
            knowledge._embedder = Embedder()
            context = knowledge._format_context(
                "automobile maintenance",
                knowledge._build_snapshot(documents, vectors),
            )

        self.assertIn("vehicle service handbook", context)
        self.assertNotIn("kitchen recipe collection", context)

    def test_follow_up_search_uses_original_and_contextual_queries(self):
        self.assertEqual(
            _query_texts(
                "And the net profit in the same third year?",
                "What is third-year EBITDA for the AulaLayer project?",
            ),
            (
                "And the net profit in the same third year",
                "What is third-year EBITDA for the AulaLayer project\n"
                "And the net profit in the same third year",
            ),
        )

    def test_contextual_query_disambiguates_follow_up(self):
        class Reranker:
            def predict(self, pairs, **kwargs):
                return [
                    10.0 if "AulaLayer" in passage else 0.0
                    for _, passage in pairs
                ]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "knowledgebase"
            root.mkdir()
            knowledge = KnowledgeBase(
                root,
                cache_path=Path(directory) / "index.sqlite3",
                watch=False,
                top_chunks=1,
                reranker_model="test",
            )
            knowledge._store.sync_document(
                "beacon.txt",
                (1, 1),
                [("Forecast", None, "Year 3 net profit: EUR 999", "Project: Beacon")],
                None,
            )
            knowledge._store.sync_document(
                "aulalayer.txt",
                (1, 1),
                [("Forecast", None, "Year 3 net profit: EUR 780", "Project: AulaLayer")],
                None,
            )
            knowledge._reranker = Reranker()

            context = knowledge._format_store_context(
                "And the net profit in the same third year?",
                "What is third-year EBITDA for the AulaLayer project?",
            )
            knowledge.stop()

        self.assertIn("EUR 780", context)
        self.assertNotIn("EUR 999", context)
        self.assertIn("do not derive or calculate", context)

    def test_agent_tools_search_filter_neighbors_and_share_budget(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "knowledgebase"
            root.mkdir()
            knowledge = KnowledgeBase(
                root,
                cache_path=Path(directory) / "index.sqlite3",
                watch=False,
            )
            knowledge._store.sync_document(
                "alpha.txt",
                (1, 1),
                [
                    ("Forecast", 1, "Year 3", "Project Alpha"),
                    ("Forecast", 2, "Net profit: EUR 780", "Project Alpha"),
                ],
                None,
            )
            knowledge._store.sync_document(
                "beta.txt",
                (1, 1),
                [("Forecast", 1, "Net profit: EUR 999", "Project Beta")],
                None,
            )
            knowledge._ready = True
            tools = {tool.name: tool for tool in knowledge.agent_tools()}
            search = tools["search_knowledge"]
            context = tools["get_knowledge_context"]
            calculator = tools["calculate_knowledge"]

            result = search.execute("net profit", source="alpha.txt", limit=2)
            chunk_id = int(result.split("[chunk: ", 1)[1].split("]", 1)[0])
            neighbors = context.execute(chunk_id, before=1, after=1)
            search.execute("tax", source="alpha.txt")
            exhausted = context.execute(chunk_id)
            calculation = calculator.execute("(780 - 700) / 700 * 100")
            rejected = calculator.run(expression="2 ** 8")
            metrics = knowledge.metrics()
            knowledge.stop()

        self.assertIn("EUR 780", result)
        self.assertNotIn("EUR 999", result)
        self.assertIn("Year 3", neighbors)
        self.assertIn("budget exhausted", exhausted)
        self.assertIn("11.428571", calculation)
        self.assertIn("[Tool Error]", rejected)
        self.assertIn("Agentic RAG calls: 3", metrics)
        self.assertEqual(search.execution_delay, 0.0)
        self.assertTrue(search.turn_scoped)

    def test_sqlite_cache_reuses_chunks_and_embeddings(self):
        class Embedder:
            def encode(self, texts, **kwargs):
                return np.array([[1.0, 0.0] for _ in texts], dtype=np.float32)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "knowledgebase"
            root.mkdir()
            document_path = root / "manual.txt"
            document_path.write_text("Cached maintenance procedure.")
            cache_path = Path(directory) / "cache" / "index.sqlite3"
            calls = []

            def extractor(path, base):
                calls.append(path)
                source = str(path.relative_to(base))
                return _Document(source, (
                    _Chunk(source, "Manual", None, path.read_text(), "Document: Manual"),
                ))

            first = KnowledgeBase(
                root,
                watch=False,
                cache_path=cache_path,
                embedding_model="test",
                extractor=extractor,
            )
            first._embedder = Embedder()
            first._refresh_sync()
            self.assertEqual(len(calls), 1)

            def fail_extractor(path, base):
                raise AssertionError("unchanged cached document was reparsed")

            second = KnowledgeBase(
                root,
                watch=False,
                cache_path=cache_path,
                embedding_model="test",
                extractor=fail_extractor,
            )
            second._embedder = Embedder()
            second._refresh_sync()
            document_count = second._store.count_documents()
            chunk_count = second._store.count_chunks()
            dense_hits = second._store.dense_search(
                np.array([1.0, 0.0], dtype=np.float32), 1
            )

        self.assertEqual(document_count, 1)
        self.assertEqual(chunk_count, 1)
        self.assertEqual(len(dense_hits), 1)

    def test_persistent_hybrid_query_records_metrics(self):
        class Embedder:
            def encode(self, texts, **kwargs):
                vectors = []
                for text in texts:
                    vectors.append(
                        [1.0, 0.0]
                        if "automobile" in text or "vehicle" in text
                        else [0.0, 1.0]
                    )
                return np.asarray(vectors, dtype=np.float32)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "knowledgebase"
            root.mkdir()
            (root / "vehicles.txt").write_text("vehicle service handbook")
            (root / "cooking.txt").write_text("bread baking recipe")
            knowledge = KnowledgeBase(
                root,
                watch=False,
                cache_path=Path(directory) / "cache" / "index.sqlite3",
                embedding_model="test",
                top_chunks=1,
            )
            knowledge._embedder = Embedder()
            knowledge._refresh_sync()
            context = knowledge._format_store_context("automobile upkeep")
            metrics = knowledge.metrics()

        self.assertIn("vehicle service handbook", context)
        self.assertIn("Queries sampled: 1", metrics)
        self.assertIn("dense_ms", metrics)

    def test_embedding_model_change_reuses_chunks_and_reembeds(self):
        class Embedder:
            def __init__(self, vector):
                self.vector = vector

            def encode(self, texts, **kwargs):
                return np.array([self.vector for _ in texts], dtype=np.float32)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "knowledgebase"
            root.mkdir()
            document_path = root / "manual.txt"
            document_path.write_text("Reusable extracted content.")
            cache_path = Path(directory) / "cache" / "index.sqlite3"
            calls = []

            def extractor(path, base):
                calls.append(path)
                source = str(path.relative_to(base))
                return _Document(source, (_Chunk(source, "Manual", None, path.read_text()),))

            first = KnowledgeBase(
                root,
                watch=False,
                cache_path=cache_path,
                embedding_model="model-a",
                extractor=extractor,
            )
            first._embedder = Embedder([1.0, 0.0])
            first._refresh_sync()

            def fail_extractor(path, base):
                raise AssertionError("model change should not reparse documents")

            second = KnowledgeBase(
                root,
                watch=False,
                cache_path=cache_path,
                embedding_model="model-b",
                extractor=fail_extractor,
            )
            second._embedder = Embedder([0.0, 1.0])
            second._refresh_sync()
            hits = second._store.dense_search(
                np.array([0.0, 1.0], dtype=np.float32),
                1,
            )

        self.assertEqual(len(calls), 1)
        self.assertEqual(len(hits), 1)
        self.assertAlmostEqual(hits[0][1], 1.0)

    def test_changed_file_invalidates_sqlite_cache(self):
        class Embedder:
            def encode(self, texts, **kwargs):
                return np.array([[1.0, 0.0] for _ in texts], dtype=np.float32)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "knowledgebase"
            root.mkdir()
            document_path = root / "manual.txt"
            document_path.write_text("First version.")
            cache_path = Path(directory) / "cache" / "index.sqlite3"
            calls = []

            def extractor(path, base):
                calls.append(path.read_text())
                source = str(path.relative_to(base))
                return _Document(source, (_Chunk(source, "Manual", None, path.read_text()),))

            first = KnowledgeBase(
                root,
                watch=False,
                cache_path=cache_path,
                embedding_model="test",
                extractor=extractor,
            )
            first._embedder = Embedder()
            first._refresh_sync()
            document_path.write_text("Second version with a different size.")

            second = KnowledgeBase(
                root,
                watch=False,
                cache_path=cache_path,
                embedding_model="test",
                extractor=extractor,
            )
            second._embedder = Embedder()
            second._refresh_sync()

        self.assertEqual(calls, ["First version.", "Second version with a different size."])

    async def test_unrelated_query_returns_no_context(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "rockets.txt").write_text("Orbital launch sequence and propellant loading.")
            knowledge = KnowledgeBase(root, watch=False)
            knowledge.start()
            await knowledge._refresh_task

            context = await knowledge.retrieve("How do I bake sourdough?")
            knowledge.stop()

        self.assertEqual(context, "")

    async def test_query_does_not_wait_for_initial_index(self):
        gate = threading.Event()

        def slow_extractor(path, root):
            gate.wait(timeout=2)
            return _Document(
                str(path.relative_to(root)),
                (_Chunk(path.name, path.stem, None, path.read_text()),),
            )

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "manual.txt").write_text("Emergency shutdown procedure.")
            knowledge = KnowledgeBase(root, watch=False, extractor=slow_extractor)
            knowledge.start()
            await asyncio.sleep(0.01)

            context = await knowledge.retrieve("shutdown")
            gate.set()
            await knowledge._refresh_task
            knowledge.stop()

        self.assertEqual(context, "")

    async def test_reindex_replaces_changed_content(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            document = root / "manual.txt"
            document.write_text("The access code is amber.")
            knowledge = KnowledgeBase(root, watch=False)
            knowledge.start()
            await knowledge._refresh_task
            self.assertIn("amber", await knowledge.retrieve("access code"))

            document.write_text("The access code is cobalt.")
            self.assertEqual(await knowledge.retrieve("access code"), "")
            await asyncio.sleep(0.5)
            await knowledge._refresh_task
            context = await knowledge.retrieve("access code")
            knowledge.stop()

        self.assertIn("cobalt", context)
        self.assertNotIn("amber", context)

    async def test_watcher_indexes_new_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            knowledge = KnowledgeBase(root, watch=True)
            knowledge.start()
            await knowledge._refresh_task
            (root / "new.txt").write_text("Newly indexed procedure.")

            for _ in range(30):
                await asyncio.sleep(0.1)
                with knowledge._lock:
                    if knowledge._snapshot.documents:
                        break
            with knowledge._lock:
                sources = [doc.source for doc in knowledge._snapshot.documents]
            knowledge.stop()

        self.assertEqual(sources, ["new.txt"])

    def test_context_metadata_disambiguates_similar_chunks(self):
        documents = [
            _Document("orion.txt", (
                _Chunk(
                    "orion.txt",
                    "Schedule",
                    None,
                    "The deadline is target-value.",
                    "Document: Orion handbook\nSection: Schedule",
                ),
            )),
            _Document("nova.txt", (
                _Chunk(
                    "nova.txt",
                    "Schedule",
                    None,
                    "The deadline is distractor-value.",
                    "Document: Nova handbook\nSection: Schedule",
                ),
            )),
        ]
        with tempfile.TemporaryDirectory() as directory:
            knowledge = KnowledgeBase(Path(directory), watch=False, top_chunks=1)
            context = knowledge._format_context(
                "What is the Orion deadline?",
                knowledge._build_snapshot(documents),
            )

        self.assertIn("target-value", context)
        self.assertNotIn("distractor-value", context)

    def test_table_rows_preserve_key_value_relationships(self):
        rows = _table_rows(
            "| Field | Year 2 | Year 3 |\n"
            "|---|---|---|\n"
            "| Net profit | 700 | 780 |\n"
        )

        self.assertTrue(any("Net profit: 700 | 780" in row for row in rows))
        self.assertTrue(any("Year 2: 700" in row for row in rows))
        self.assertTrue(any("Year 3: 780" in row for row in rows))

    def test_repeated_table_title_becomes_structural_heading(self):
        markdown = (
            "| Quarterly Performance Report - Metric | Quarterly Performance Report - Value |\n"
            "|---|---|\n"
            "| Revenue | 10 |\n"
        )

        self.assertEqual(_table_heading(markdown), "Quarterly Performance Report")

    async def test_markdown_headings_and_tables_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            document = root / "vehicle.md"
            document.write_text(
                "# Vehicle Guide\n\n## Launch Sequence\n\n"
                "| Step | Action |\n|---|---|\n| 1 | Load propellant |\n"
            )
            knowledge = KnowledgeBase(root, watch=False)
            indexed = await asyncio.to_thread(knowledge._extract_document, document, root)
            context = knowledge._format_context(
                "launch sequence",
                knowledge._build_snapshot([indexed]),
            )

        self.assertTrue(any("Launch Sequence" in chunk.section for chunk in indexed.chunks))
        self.assertTrue(any("Load propellant" in chunk.text for chunk in indexed.chunks))
        self.assertIn("Load propellant", context)

    async def test_pptx_slides_are_discovered_and_extracted(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            document = root / "forecast.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "AulaLayer Financial Forecast"
            slide.placeholders[1].text = "Year 3 net profit is EUR 780,000."
            table = slide.shapes.add_table(
                2, 2, Inches(1), Inches(4), Inches(6), Inches(1)
            ).table
            table.cell(0, 0).text = "Metric"
            table.cell(0, 1).text = "Year 3"
            table.cell(1, 0).text = "Revenue"
            table.cell(1, 1).text = "1,250,000"
            presentation.save(document)
            knowledge = KnowledgeBase(root, watch=False)
            discovered = document in knowledge._files()
            indexed = await asyncio.to_thread(knowledge._extract_document, document, root)

        self.assertTrue(discovered)
        self.assertEqual(indexed.source, "forecast.pptx")
        self.assertTrue(any("AulaLayer" in chunk.section for chunk in indexed.chunks))
        self.assertTrue(any("780,000" in chunk.text for chunk in indexed.chunks))
        self.assertTrue(any("Revenue" in chunk.text for chunk in indexed.chunks))
        self.assertTrue(any("Year 3: 1,250,000" in chunk.text for chunk in indexed.chunks))


class RAGCommandTests(unittest.TestCase):
    def test_status_and_reindex_commands(self):
        class RAG:
            def __init__(self):
                self.reindexed = False

            def status(self):
                return "RAG [ready]"

            def metrics(self):
                return "RAG metrics"

            def request_refresh(self, force=False):
                self.reindexed = force
                return True

        rag = RAG()
        core = AndrewCore()
        core.domain = SimpleNamespace(rag=rag)

        self.assertEqual(core.handle_slash("/rag status", None), "RAG [ready]")
        self.assertEqual(core.handle_slash("/rag metrics", None), "RAG metrics")
        self.assertEqual(
            core.handle_slash("/rag reindex", None),
            "✓ Knowledge base reindex started.",
        )
        self.assertTrue(rag.reindexed)


class TemporaryContextTests(unittest.TestCase):
    def test_turn_scoped_tool_evidence_is_removed_from_memory(self):
        llm = LLM.__new__(LLM)
        llm.memory = Memory.__new__(Memory)
        llm.memory.messages = [
            {"role": "user", "content": "Find the value"},
            {
                "role": "assistant",
                "content": "",
                "tool_calls": [
                    {
                        "id": "rag-call",
                        "function": {"name": "search_knowledge", "arguments": "{}"},
                    },
                ],
            },
            {
                "role": "tool",
                "tool_call_id": "rag-call",
                "content": "Large retrieved evidence",
            },
            {"role": "assistant", "content": "The cited final answer"},
        ]
        tool = SimpleNamespace(name="search_knowledge", turn_scoped=True)

        llm._strip_turn_scoped_tools(0, [tool])

        self.assertEqual(
            llm.memory.messages,
            [
                {"role": "user", "content": "Find the value"},
                {"role": "assistant", "content": "The cited final answer"},
            ],
        )

    def test_last_user_prompt_excludes_assistant_content(self):
        memory = Memory.__new__(Memory)
        memory.messages = [
            {"role": "user", "content": "Question about AulaLayer"},
            {"role": "assistant", "content": "An unsupported inferred value"},
        ]

        self.assertEqual(memory.last_user_prompt, "Question about AulaLayer")

    def test_knowledge_context_is_not_stored_in_memory(self):
        memory = Memory.__new__(Memory)
        memory.system_prompt = "Base prompt"
        memory.messages = [{"role": "user", "content": "Question"}]
        memory.summary = ""
        memory.last_exchange = ""
        memory._trimmed = False
        memory._active_skills = []

        messages = memory.get("<knowledge>Temporary fact</knowledge>")

        self.assertIn("Temporary fact", messages[0]["content"])
        self.assertNotIn("Temporary fact", str(memory.messages))


if __name__ == "__main__":
    unittest.main()
